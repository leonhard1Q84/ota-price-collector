# -*- coding: utf-8 -*-
"""Build and write Klook BM weekly reports from local BM files."""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_CONFIG = SCRIPT_DIR / "bm_weekly_report_config.json"
EXAMPLE_CONFIG = SCRIPT_DIR / "bm_weekly_report_config.example.json"

BM_FILENAME_RE = re.compile(
    r"^(?P<code>[A-Z0-9]+)_BM_?(?P<date>\d{8})_(?P<time>\d{6})\.(?P<ext>csv|xlsx)$",
    re.IGNORECASE,
)
ABP_LOR_RE = re.compile(r"ABP=(?P<abp>\d+).*?LOR=(?P<lor>\d+)", re.IGNORECASE)
REPORT_BUCKETS = ("1-7", "8-14", "15-30", "30-60")
REPORT_LORS = (2, 4, 6)
REPORT_KEYS = tuple((bucket, lor) for bucket in REPORT_BUCKETS for lor in REPORT_LORS)
WEEK_BLOCK_WIDTH = 12


class BMReportError(RuntimeError):
    pass


@dataclass(frozen=True)
class BMFile:
    path: Path
    code: str
    file_date: date
    time_text: str
    ext: str


@dataclass(frozen=True)
class SheetBlock:
    header_row_idx: int
    header_row_number: int
    row_by_name: Dict[str, int]


@dataclass(frozen=True)
class SheetMatrix:
    row_indices: List[int]
    col_indices: List[str]
    values: List[List[str]]


def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="汇总 Klook BM 文件并写入 BM 周均值报表")
    parser.add_argument("--config", help="配置文件路径，默认 src/bm_weekly_report_config.json")
    parser.add_argument("--run-date", help="执行日期 YYYY-MM-DD；默认今天。周五执行时统计上周五到本周四")
    parser.add_argument("--week-start", help="手动指定统计开始日期 YYYY-MM-DD")
    parser.add_argument("--week-end", help="手动指定统计结束日期 YYYY-MM-DD")
    parser.add_argument("--week-label", help="手动指定周标签，例如 W27；默认使用统计结束日期的 ISO 周")
    parser.add_argument("--dry-run", action="store_true", help="只生成本地文件和打印待写入内容，不写飞书")
    parser.add_argument("--allow-empty", action="store_true", help="没有解析到 BM 数据时仍允许写入 '-'")
    parser.add_argument("--no-insert", action="store_true", help="目标周不存在时不插入新周块，只报错")
    parser.add_argument("--write-feishu", action="store_true", help="正式写入飞书；未传时等同 dry-run")
    return parser.parse_args(argv)


def parse_date(value: str) -> date:
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError as exc:
        raise BMReportError(f"日期格式错误，应为 YYYY-MM-DD：{value}") from exc


def default_week_window(run_day: date) -> Tuple[date, date]:
    days_since_thursday = (run_day.weekday() - 3) % 7
    end_day = run_day - timedelta(days=days_since_thursday)
    if run_day.weekday() == 4:
        end_day = run_day - timedelta(days=1)
    return end_day - timedelta(days=6), end_day


def iso_weeks_in_year(year: int) -> int:
    return date(year, 12, 28).isocalendar().week


def week_label_for(reference_day: date, offset: int = 0) -> str:
    week = reference_day.isocalendar().week + int(offset)
    if week <= 0:
        week += iso_weeks_in_year(start_day.year - 1)
    return f"W{week}"


def load_config(config_path: Optional[str]) -> Dict[str, Any]:
    path = Path(config_path) if config_path else DEFAULT_CONFIG
    if not path.exists():
        if not config_path:
            raise BMReportError(
                f"找不到正式配置文件：{path}；请先复制 {EXAMPLE_CONFIG.name} 为 {DEFAULT_CONFIG.name} 后再运行"
            )
        raise BMReportError(f"找不到配置文件：{path}")
    with path.open("r", encoding="utf-8") as f:
        config = json.load(f)
    config["_config_path"] = str(path)
    return config


def is_windows_absolute_path(text: str) -> bool:
    return bool(re.match(r"^[A-Za-z]:[\\/]", text)) or text.startswith("\\\\")


def resolve_path(value: str) -> Path:
    text = os.path.expandvars(str(value))
    if os.name != "nt" and is_windows_absolute_path(text):
        return Path(text)
    if os.name != "nt":
        text = text.replace("\\", os.sep)
    path = Path(text)
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    return path


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def station_alias_candidates(text: str) -> List[str]:
    raw = normalize_text(text)
    candidates: List[str] = []

    def add(value: str) -> None:
        value = normalize_text(value)
        if value and value not in candidates:
            candidates.append(value)

    add(raw)
    add(raw.replace("+", " "))
    add(raw.replace("+", ""))
    no_paren = re.sub(r"\s*[（(][^）)]*[）)]\s*", "", raw).strip()
    add(no_paren)
    add(no_paren.replace("+", " "))
    add(no_paren.replace("+", ""))
    return candidates


def normalize_station(raw: str, config: Dict[str, Any], code: str = "") -> str:
    text = normalize_text(raw)
    aliases = config.get("station_aliases") or {}
    for candidate in station_alias_candidates(text):
        if candidate in aliases:
            return normalize_text(aliases[candidate])
    if not text:
        defaults = config.get("code_default_stations") or {}
        return normalize_text(defaults.get(code.upper(), ""))
    return station_alias_candidates(text)[-1] if station_alias_candidates(text) else text


def bucket_for_abp(abp: int) -> Optional[str]:
    if 0 <= abp <= 6:
        return "1-7"
    if 7 <= abp <= 14:
        return "8-14"
    if 15 <= abp <= 30:
        return "15-30"
    if 31 <= abp <= 60:
        return "30-60"
    return None


def parse_percent(value: Any) -> Optional[float]:
    text = normalize_text(value)
    if not text or text in {"-", "--"} or text.lower() in {"nan", "none", "null"}:
        return None
    text = text.replace(",", "")
    if text.endswith("%"):
        text = text[:-1].strip()
        try:
            return float(text) / 100.0
        except ValueError:
            return None
    try:
        number = float(text)
    except ValueError:
        return None
    if abs(number) > 1:
        return number / 100.0
    return number


def format_percent(value: Optional[float], decimals: int = 0) -> str:
    if value is None:
        return "-"
    scale = Decimal("1") if decimals <= 0 else Decimal("1").scaleb(-decimals)
    pct = (Decimal(str(value)) * Decimal("100")).quantize(scale, rounding=ROUND_HALF_UP)
    if decimals <= 0:
        return f"{int(pct)}%"
    return f"{pct:.{decimals}f}%"


def average(values: Iterable[float]) -> Optional[float]:
    items = [v for v in values if v is not None]
    if not items:
        return None
    return sum(items) / len(items)


def bm_file_rank(item: BMFile) -> Tuple[str, int, str]:
    try:
        mtime_ns = item.path.stat().st_mtime_ns
    except OSError:
        mtime_ns = 0
    return item.time_text, mtime_ns, item.path.name.lower()


def collect_bm_files(config: Dict[str, Any], start_day: date, end_day: date) -> List[BMFile]:
    files: Dict[Tuple[str, date], BMFile] = {}
    for root_text in config.get("data_roots") or []:
        root = resolve_path(root_text)
        if not root.exists():
            print(f"跳过不存在的数据目录：{root}")
            continue
        for path in root.rglob("*"):
            if not path.is_file():
                continue
            match = BM_FILENAME_RE.match(path.name)
            if not match:
                continue
            file_day = datetime.strptime(match.group("date"), "%Y%m%d").date()
            if not (start_day <= file_day <= end_day):
                continue
            candidate = BMFile(
                path=path,
                code=match.group("code").upper(),
                file_date=file_day,
                time_text=match.group("time"),
                ext=match.group("ext").lower(),
            )
            key = (candidate.code, candidate.file_date)
            existing = files.get(key)
            if existing is None or bm_file_rank(candidate) > bm_file_rank(existing):
                files[key] = candidate
    return sorted(files.values(), key=lambda item: (item.file_date, item.time_text, item.path.name))


def read_csv_rows(path: Path) -> List[List[str]]:
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            with path.open("r", encoding=encoding, newline="") as f:
                return [[normalize_text(cell) for cell in row] for row in csv.reader(f)]
        except UnicodeDecodeError:
            continue
    raise BMReportError(f"无法读取 CSV 编码：{path}")


def read_xlsx_rows(path: Path) -> List[List[str]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise BMReportError("缺少 openpyxl 依赖，请运行：python -m pip install -r requirements.txt") from exc
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["劣势分析"] if "劣势分析" in workbook.sheetnames else workbook.active
    rows: List[List[str]] = []
    for row in sheet.iter_rows(values_only=True):
        rows.append([normalize_text(cell) for cell in row])
    workbook.close()
    return rows


def read_table_rows(path: Path) -> List[List[str]]:
    if path.suffix.lower() == ".xlsx":
        return read_xlsx_rows(path)
    return read_csv_rows(path)


def find_header_row(rows: List[List[str]]) -> Optional[int]:
    for idx, row in enumerate(rows):
        if any(ABP_LOR_RE.search(normalize_text(cell)) for cell in row):
            return idx
    return None


def find_column(header: List[str], names: Iterable[str]) -> Optional[int]:
    wanted = {normalize_text(name) for name in names}
    for idx, cell in enumerate(header):
        if normalize_text(cell) in wanted:
            return idx
    return None


def parse_bm_file(item: BMFile, config: Dict[str, Any]) -> Dict[str, Dict[Tuple[str, int], List[float]]]:
    rows = read_table_rows(item.path)
    header_idx = find_header_row(rows)
    if header_idx is None:
        print(f"WARNING 未找到 ABP/LOR 表头，跳过：{item.path}")
        return {}
    header = rows[header_idx]
    pickup_col = find_column(header, ("取车点", "门店", "站点"))
    if pickup_col is None and header_idx > 0:
        pickup_col = find_column(rows[header_idx - 1], ("取车点", "门店", "站点"))

    metric_columns: Dict[int, Tuple[str, int]] = {}
    for col_idx, cell in enumerate(header):
        match = ABP_LOR_RE.search(normalize_text(cell))
        if not match:
            continue
        abp = int(match.group("abp"))
        lor = int(match.group("lor"))
        bucket = bucket_for_abp(abp)
        if bucket and lor in REPORT_LORS:
            metric_columns[col_idx] = (bucket, lor)

    if not metric_columns:
        print(f"WARNING 未找到可统计的 ABP/LOR 列，跳过：{item.path}")
        return {}

    results: Dict[str, Dict[Tuple[str, int], List[float]]] = {}
    for row in rows[header_idx + 1 :]:
        row_name = normalize_text(row[0] if row else "")
        if row_name != "BM%":
            continue
        raw_station = ""
        if pickup_col is not None and pickup_col < len(row):
            raw_station = row[pickup_col]
        station = normalize_station(raw_station, config, item.code)
        if not station:
            print(f"WARNING BM 行没有取车点且无默认站点，跳过：{item.path}")
            continue
        station_values = results.setdefault(station, {})
        for col_idx, key in metric_columns.items():
            value = parse_percent(row[col_idx] if col_idx < len(row) else "")
            if value is None:
                continue
            station_values.setdefault(key, []).append(value)
    return results


def merge_station_samples(target: Dict[str, Dict[Tuple[str, int], List[float]]], source: Dict[str, Dict[Tuple[str, int], List[float]]]) -> None:
    for station, metrics in source.items():
        dest_metrics = target.setdefault(station, {})
        for key, values in metrics.items():
            dest_metrics.setdefault(key, []).extend(values)


def aggregate_station_values(samples: Dict[str, Dict[Tuple[str, int], List[float]]]) -> Dict[str, Dict[Tuple[str, int], Optional[float]]]:
    aggregated: Dict[str, Dict[Tuple[str, int], Optional[float]]] = {}
    for station, metrics in samples.items():
        aggregated[station] = {key: average(metrics.get(key, [])) for key in REPORT_KEYS}
    return aggregated


def aggregate_groups(
    station_values: Dict[str, Dict[Tuple[str, int], Optional[float]]],
    groups: Dict[str, List[str]],
) -> Dict[str, Dict[Tuple[str, int], Optional[float]]]:
    result: Dict[str, Dict[Tuple[str, int], Optional[float]]] = {}
    for group_name, stations in groups.items():
        group_metrics: Dict[Tuple[str, int], Optional[float]] = {}
        for key in REPORT_KEYS:
            values = [
                station_values.get(station, {}).get(key)
                for station in stations
                if station_values.get(station, {}).get(key) is not None
            ]
            group_metrics[key] = average(values)
        result[group_name] = group_metrics
    return result


def row_values(metrics: Dict[Tuple[str, int], Optional[float]], decimals: int) -> List[str]:
    return [format_percent(metrics.get(key), decimals) for key in REPORT_KEYS]


def write_tsv(path: Path, rows: List[List[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter="\t", lineterminator="\n")
        writer.writerows(rows)


def build_output_rows(names: List[str], values: Dict[str, Dict[Tuple[str, int], Optional[float]]], decimals: int) -> List[List[str]]:
    header = ["名称"] + [f"{bucket}_LOR{lor}" for bucket, lor in REPORT_KEYS]
    rows = [header]
    for name in names:
        rows.append([name] + row_values(values.get(name, {}), decimals))
    return rows


def col_to_index(col: str) -> int:
    result = 0
    for char in col.upper():
        if not ("A" <= char <= "Z"):
            raise BMReportError(f"非法列名：{col}")
        result = result * 26 + ord(char) - ord("A") + 1
    return result


def index_to_col(index: int) -> str:
    if index <= 0:
        raise BMReportError(f"非法列号：{index}")
    chars = []
    while index:
        index, rem = divmod(index - 1, 26)
        chars.append(chr(ord("A") + rem))
    return "".join(reversed(chars))


def resolve_lark_cli() -> str:
    cli = shutil.which("lark-cli")
    if cli:
        return cli
    fallback = Path.home() / ".local" / "node-v22.22.2" / "bin" / "lark-cli"
    if fallback.exists():
        return str(fallback)
    raise BMReportError("找不到 lark-cli，无法写入飞书表格")


def run_lark(command: List[str]) -> subprocess.CompletedProcess[str]:
    normalized = list(command)
    if normalized and normalized[0] == "lark-cli":
        normalized[0] = resolve_lark_cli()
    return subprocess.run(
        normalized,
        text=True,
        capture_output=True,
        cwd=str(PROJECT_ROOT),
        encoding="utf-8",
        errors="replace",
    )


def parse_json_output(stdout: str) -> Dict[str, Any]:
    text = (stdout or "").strip()
    start = text.find("{")
    if start > 0:
        text = text[start:]
    if not text:
        raise BMReportError("飞书命令没有返回 JSON")
    data = json.loads(text)
    if not isinstance(data, dict):
        raise BMReportError(f"飞书命令返回异常：{data}")
    return data


def run_lark_json(command: List[str]) -> Dict[str, Any]:
    result = run_lark(command)
    if result.returncode != 0:
        raise BMReportError(result.stderr.strip() or result.stdout.strip() or "飞书命令执行失败")
    data = parse_json_output(result.stdout)
    if data.get("ok") is False:
        raise BMReportError(json.dumps(data, ensure_ascii=False)[:1000])
    return data


def read_sheet(token: str, sheet_id: str, range_name: str) -> SheetMatrix:
    data = run_lark_json(
        [
            "lark-cli",
            "sheets",
            "+cells-get",
            "--as",
            "bot",
            "--spreadsheet-token",
            token,
            "--sheet-id",
            sheet_id,
            "--range",
            range_name,
            "--include",
            "value",
            "--format",
            "json",
        ]
    )
    ranges = ((data.get("data") or {}).get("ranges") or [])
    if not ranges:
        raise BMReportError(f"读取飞书范围为空：{sheet_id}!{range_name}")
    item = ranges[0]
    row_indices = [int(v) for v in item.get("row_indices") or []]
    col_indices = [str(v) for v in item.get("col_indices") or []]
    raw_cells = item.get("cells") or []
    values: List[List[str]] = []
    width = len(col_indices)
    for row in raw_cells:
        values.append(([normalize_text(cell.get("value")) if isinstance(cell, dict) else "" for cell in row] + [""] * width)[:width])
    return SheetMatrix(row_indices=row_indices, col_indices=col_indices, values=values)


def find_sheet_blocks(matrix: SheetMatrix) -> List[SheetBlock]:
    blocks: List[SheetBlock] = []
    for row_idx, row in enumerate(matrix.values):
        if not row or normalize_text(row[0]) != "BM周均值":
            continue
        row_by_name: Dict[str, int] = {}
        data_idx = row_idx + 3
        while data_idx < len(matrix.values):
            first = normalize_text(matrix.values[data_idx][0] if matrix.values[data_idx] else "")
            if not first:
                break
            if first == "BM周均值":
                break
            row_by_name[first] = matrix.row_indices[data_idx]
            data_idx += 1
        blocks.append(
            SheetBlock(
                header_row_idx=row_idx,
                header_row_number=matrix.row_indices[row_idx],
                row_by_name=row_by_name,
            )
        )
    return blocks


def find_week_start_col(matrix: SheetMatrix, block: SheetBlock, week_label: str) -> Optional[str]:
    header = matrix.values[block.header_row_idx]
    for col_idx in range(1, len(header)):
        if normalize_text(header[col_idx]) == week_label:
            return matrix.col_indices[col_idx]
    return None


def write_sheet_range(token: str, sheet_id: str, cell_range: str, values: List[List[str]], label: str) -> None:
    cells = [[{"value": value} for value in row] for row in values]
    result = run_lark(
        [
            "lark-cli",
            "sheets",
            "+cells-set",
            "--as",
            "bot",
            "--spreadsheet-token",
            token,
            "--sheet-id",
            sheet_id,
            "--range",
            cell_range,
            "--cells",
            json.dumps(cells, ensure_ascii=False),
        ]
    )
    if result.returncode != 0:
        raise BMReportError(f"{label}写入失败 {cell_range}：{result.stderr.strip() or result.stdout.strip()}")
    print(f"{label}写入成功：{sheet_id}!{cell_range}")


def lark_exec(command: List[str], label: str) -> None:
    result = run_lark(command)
    if result.returncode != 0:
        raise BMReportError(f"{label}失败：{result.stderr.strip() or result.stdout.strip()}")
    print(f"{label}成功")


def merge_range(token: str, sheet_id: str, cell_range: str) -> None:
    lark_exec(
        [
            "lark-cli",
            "sheets",
            "+cells-merge",
            "--as",
            "bot",
            "--spreadsheet-token",
            token,
            "--sheet-id",
            sheet_id,
            "--range",
            cell_range,
        ],
        f"合并 {sheet_id}!{cell_range}",
    )


def unmerge_range(token: str, sheet_id: str, cell_range: str) -> None:
    lark_exec(
        [
            "lark-cli",
            "sheets",
            "+cells-unmerge",
            "--as",
            "bot",
            "--spreadsheet-token",
            token,
            "--sheet-id",
            sheet_id,
            "--range",
            cell_range,
        ],
        f"取消合并 {sheet_id}!{cell_range}",
    )


def unmerge_week_header_area(token: str, sheet_id: str, start_col: str, block: SheetBlock, width: int = WEEK_BLOCK_WIDTH) -> None:
    start_index = col_to_index(start_col)
    end_col = index_to_col(start_index + width - 1)
    start_row = block.header_row_number
    unmerge_range(token, sheet_id, f"{start_col}{start_row}:{end_col}{start_row + 2}")


def header_rows(week_label: str) -> List[List[str]]:
    return [
        [week_label] + [""] * 11,
        ["1-7", "", "", "8-14", "", "", "15-30", "", "", "30-60", "", ""],
        [str(lor) for _bucket in REPORT_BUCKETS for lor in REPORT_LORS],
    ]


def write_block_headers(token: str, sheet_id: str, start_col: str, block: SheetBlock, week_label: str) -> None:
    start_index = col_to_index(start_col)
    end_col = index_to_col(start_index + 11)
    start_row = block.header_row_number
    write_sheet_range(
        token,
        sheet_id,
        f"{start_col}{start_row}:{end_col}{start_row + 2}",
        header_rows(week_label),
        "BM周报表头",
    )


def merge_block_headers(token: str, sheet_id: str, start_col: str, block: SheetBlock) -> None:
    start_index = col_to_index(start_col)
    start_row = block.header_row_number
    merge_range(token, sheet_id, f"{start_col}{start_row}:{index_to_col(start_index + WEEK_BLOCK_WIDTH - 1)}{start_row}")
    for offset in (0, 3, 6, 9):
        left = index_to_col(start_index + offset)
        right = index_to_col(start_index + offset + 2)
        merge_range(token, sheet_id, f"{left}{start_row + 1}:{right}{start_row + 1}")


def matrix_value(matrix: SheetMatrix, row_number: int, col: str) -> str:
    try:
        row_idx = matrix.row_indices.index(row_number)
        col_idx = matrix.col_indices.index(col)
    except ValueError:
        return ""
    row = matrix.values[row_idx]
    return normalize_text(row[col_idx] if col_idx < len(row) else "")


def matrix_range_values(matrix: SheetMatrix, start_col: str, end_col: str, start_row: int, end_row: int) -> List[List[str]]:
    start_idx = col_to_index(start_col)
    end_idx = col_to_index(end_col)
    rows: List[List[str]] = []
    for row_number in range(start_row, end_row + 1):
        rows.append([matrix_value(matrix, row_number, index_to_col(col_idx)) for col_idx in range(start_idx, end_idx + 1)])
    return rows


def week_label_positions(matrix: SheetMatrix, block: SheetBlock) -> List[Tuple[str, str]]:
    header = matrix.values[block.header_row_idx]
    positions: List[Tuple[str, str]] = []
    for idx, cell in enumerate(header):
        text = normalize_text(cell)
        if re.fullmatch(r"W\d+", text):
            positions.append((matrix.col_indices[idx], text))
    return positions


def detect_partial_left_insert(matrix: SheetMatrix, blocks: List[SheetBlock]) -> Optional[int]:
    if not blocks:
        return None
    positions = week_label_positions(matrix, blocks[0])
    if len(positions) < 2 or positions[0][0] != "B":
        return None
    gap = col_to_index(positions[1][0]) - col_to_index(positions[0][0])
    if gap <= WEEK_BLOCK_WIDTH or gap % WEEK_BLOCK_WIDTH != 0:
        return None
    inserted_width = gap - WEEK_BLOCK_WIDTH
    return inserted_width if inserted_width > 0 else None


def block_last_row(block: SheetBlock) -> int:
    return max([block.header_row_number + 2] + list(block.row_by_name.values()))


def repair_partial_left_insert(
    token: str,
    sheet_id: str,
    matrix: SheetMatrix,
    blocks: List[SheetBlock],
    inserted_width: int,
) -> None:
    old_start_col = index_to_col(col_to_index("B") + inserted_width)
    for block in blocks:
        start_row = block.header_row_number
        last_row = block_last_row(block)
        expanded_end_col = index_to_col(col_to_index("B") + inserted_width + WEEK_BLOCK_WIDTH - 1)
        unmerge_range(token, sheet_id, f"B{start_row}:{expanded_end_col}{start_row + 2}")
        values = [[matrix_value(matrix, row_number, "B")] for row_number in range(start_row, last_row + 1)]
        write_sheet_range(
            token,
            sheet_id,
            f"{old_start_col}{start_row}:{old_start_col}{last_row}",
            values,
            "恢复已插入周列",
        )
        merge_block_headers(token, sheet_id, old_start_col, block)


def insert_week_columns_with_preserved_left_block(
    token: str,
    sheet_id: str,
    matrix: SheetMatrix,
    blocks: List[SheetBlock],
) -> None:
    preserved: List[Tuple[SheetBlock, int, int, List[List[str]]]] = []
    for block in blocks:
        start_row = block.header_row_number
        last_row = block_last_row(block)
        preserved.append((block, start_row, last_row, matrix_range_values(matrix, "B", "M", start_row, last_row)))
        unmerge_week_header_area(token, sheet_id, "B", block)

    lark_exec(
        [
            "lark-cli",
            "sheets",
            "+dim-insert",
            "--as",
            "bot",
            "--spreadsheet-token",
            token,
            "--sheet-id",
            sheet_id,
            "--position",
            "B",
            "--count",
            str(WEEK_BLOCK_WIDTH),
            "--inherit-style",
            "after",
        ],
        f"插入新周列 {sheet_id}",
    )

    old_start_col = index_to_col(col_to_index("B") + WEEK_BLOCK_WIDTH)
    old_end_col = index_to_col(col_to_index(old_start_col) + WEEK_BLOCK_WIDTH - 1)
    for block, start_row, last_row, values in preserved:
        write_sheet_range(
            token,
            sheet_id,
            f"{old_start_col}{start_row}:{old_end_col}{last_row}",
            values,
            "恢复旧周数据",
        )
        merge_block_headers(token, sheet_id, old_start_col, block)


def write_report_rows(
    token: str,
    sheet_id: str,
    start_col: str,
    block: SheetBlock,
    ordered_names: List[str],
    values: Dict[str, Dict[Tuple[str, int], Optional[float]]],
    decimals: int,
    label: str,
) -> List[str]:
    start_index = col_to_index(start_col)
    end_col = index_to_col(start_index + 11)
    missing_rows: List[str] = []
    for name in ordered_names:
        row_number = block.row_by_name.get(name)
        if not row_number:
            missing_rows.append(name)
            continue
        write_sheet_range(
            token,
            sheet_id,
            f"{start_col}{row_number}:{end_col}{row_number}",
            [row_values(values.get(name, {}), decimals)],
            f"{label} {name}",
        )
    return missing_rows


def prepare_sheet_for_week(
    token: str,
    sheet_id: str,
    read_range: str,
    week_label: str,
    format_rows: int,
    no_insert: bool,
    dry_run: bool,
) -> Tuple[str, List[SheetBlock], bool]:
    matrix = read_sheet(token, sheet_id, read_range)
    blocks = find_sheet_blocks(matrix)
    if not blocks:
        raise BMReportError(f"Sheet {sheet_id} 未定位到 BM周均值 表格")
    existing_col = find_week_start_col(matrix, blocks[0], week_label)
    if existing_col:
        return existing_col, blocks, False
    partial_insert_width = detect_partial_left_insert(matrix, blocks)
    if partial_insert_width:
        if no_insert:
            raise BMReportError(f"Sheet {sheet_id} 检测到未完成的新周列，且指定了 --no-insert")
        if dry_run:
            return "B", blocks, True
        repair_partial_left_insert(token, sheet_id, matrix, blocks, partial_insert_width)
        matrix = read_sheet(token, sheet_id, read_range)
        repaired_blocks = find_sheet_blocks(matrix)
        if not repaired_blocks:
            raise BMReportError(f"Sheet {sheet_id} 修复未完成新周列后未定位到 BM周均值 表格")
        return "B", repaired_blocks, True
    if no_insert:
        raise BMReportError(f"Sheet {sheet_id} 未找到 {week_label}，且指定了 --no-insert")
    if dry_run:
        return "B", blocks, True
    insert_week_columns_with_preserved_left_block(token, sheet_id, matrix, blocks)
    matrix = read_sheet(token, sheet_id, read_range)
    inserted_blocks = find_sheet_blocks(matrix)
    if not inserted_blocks:
        raise BMReportError(f"Sheet {sheet_id} 插入新周列后未定位到 BM周均值 表格")
    return "B", inserted_blocks, True


def main(argv: Optional[List[str]] = None) -> int:
    args = parse_args(argv)
    config = load_config(args.config)
    run_day = parse_date(args.run_date) if args.run_date else date.today()
    if args.week_start or args.week_end:
        if not args.week_start or not args.week_end:
            raise BMReportError("--week-start 和 --week-end 必须同时提供")
        week_start = parse_date(args.week_start)
        week_end = parse_date(args.week_end)
    else:
        week_start, week_end = default_week_window(run_day)
    if week_end < week_start:
        raise BMReportError("统计结束日期不能早于开始日期")
    week_label = args.week_label or week_label_for(week_end, int(config.get("week_label_offset", 0)))
    decimals = int(config.get("percent_decimals", 0))
    dry_run = args.dry_run or not args.write_feishu

    print("=" * 80)
    print("Klook BM 周均值汇总")
    print(f"配置文件: {config.get('_config_path')}")
    print(f"运行日期: {run_day}")
    print(f"统计周期: {week_start} -> {week_end}")
    print(f"周标签: {week_label}")
    print(f"模式: {'dry-run，不写飞书' if dry_run else '正式写入飞书'}")
    print("=" * 80)

    files = collect_bm_files(config, week_start, week_end)
    print(f"找到 BM 文件: {len(files)} 个")
    for item in files[:30]:
        print(f"  - {item.path}")
    if len(files) > 30:
        print(f"  ... 还有 {len(files) - 30} 个")

    samples: Dict[str, Dict[Tuple[str, int], List[float]]] = {}
    for item in files:
        try:
            merge_station_samples(samples, parse_bm_file(item, config))
        except Exception as exc:
            print(f"WARNING 解析失败，已跳过：{item.path}，原因：{exc}")

    if not samples and not args.allow_empty:
        raise BMReportError("没有解析到任何 BM 数据；如需强制生成空报表，请加 --allow-empty")

    station_values = aggregate_station_values(samples)
    region_values = aggregate_groups(station_values, config.get("klook_region_groups") or {})
    city_values = aggregate_groups(station_values, config.get("city_groups") or {})

    output_dir = resolve_path(config.get("output_dir", "src\\bm_weekly_report_output"))
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    station_names = list(config.get("klook_station_rows") or [])
    region_names = list((config.get("klook_region_groups") or {}).keys())
    city_names = list((config.get("city_groups") or {}).keys())

    station_path = output_dir / f"Klook_BM周均值_{week_label}_{week_start:%Y%m%d}_{week_end:%Y%m%d}_{stamp}.tsv"
    region_path = output_dir / f"Klook_BM地区汇总_{week_label}_{week_start:%Y%m%d}_{week_end:%Y%m%d}_{stamp}.tsv"
    city_path = output_dir / f"城市BM周均值_{week_label}_{week_start:%Y%m%d}_{week_end:%Y%m%d}_{stamp}.tsv"
    write_tsv(station_path, build_output_rows(station_names, station_values, decimals))
    write_tsv(region_path, build_output_rows(region_names, region_values, decimals))
    write_tsv(city_path, build_output_rows(city_names, city_values, decimals))
    print(f"本地 Klook 站点结果: {station_path}")
    print(f"本地 Klook 地区结果: {region_path}")
    print(f"本地 城市BM 结果: {city_path}")

    discovered = sorted(samples.keys())
    print("解析到的站点: " + (", ".join(discovered) if discovered else "无"))

    if dry_run:
        print("dry-run：未写入飞书。")
        return 0
    if not samples and args.allow_empty:
        raise BMReportError("正式写入时没有解析到 BM 数据，已阻止写入空报表")

    token = normalize_text(config.get("spreadsheet_token"))
    if not token:
        raise BMReportError("配置缺少 spreadsheet_token")
    sheets = config.get("sheets") or {}
    klook_sheet = sheets.get("klook") or {}
    city_sheet = sheets.get("city") or {}

    klook_col, klook_blocks, klook_inserted = prepare_sheet_for_week(
        token,
        normalize_text(klook_sheet.get("sheet_id")),
        normalize_text(klook_sheet.get("read_range") or "A1:FI120"),
        week_label,
        int(klook_sheet.get("format_copy_rows") or 120),
        args.no_insert,
        dry_run=False,
    )
    for block in klook_blocks:
        write_block_headers(token, normalize_text(klook_sheet.get("sheet_id")), klook_col, block, week_label)
    if klook_inserted:
        for block in klook_blocks:
            merge_block_headers(token, normalize_text(klook_sheet.get("sheet_id")), klook_col, block)
    missing_station = write_report_rows(
        token,
        normalize_text(klook_sheet.get("sheet_id")),
        klook_col,
        klook_blocks[0],
        station_names,
        station_values,
        decimals,
        "Klook站点BM",
    )
    missing_region: List[str] = []
    if len(klook_blocks) > 1:
        missing_region = write_report_rows(
            token,
            normalize_text(klook_sheet.get("sheet_id")),
            klook_col,
            klook_blocks[1],
            region_names,
            region_values,
            decimals,
            "Klook地区BM",
        )

    city_col, city_blocks, city_inserted = prepare_sheet_for_week(
        token,
        normalize_text(city_sheet.get("sheet_id")),
        normalize_text(city_sheet.get("read_range") or "A1:CK80"),
        week_label,
        int(city_sheet.get("format_copy_rows") or 80),
        args.no_insert,
        dry_run=False,
    )
    for block in city_blocks:
        write_block_headers(token, normalize_text(city_sheet.get("sheet_id")), city_col, block, week_label)
    if city_inserted:
        for block in city_blocks:
            merge_block_headers(token, normalize_text(city_sheet.get("sheet_id")), city_col, block)
    missing_city = write_report_rows(
        token,
        normalize_text(city_sheet.get("sheet_id")),
        city_col,
        city_blocks[0],
        city_names,
        city_values,
        decimals,
        "城市BM",
    )

    if missing_station:
        print("WARNING Klook站点 Sheet 缺少行: " + ", ".join(missing_station))
    if missing_region:
        print("WARNING Klook地区 Sheet 缺少行: " + ", ".join(missing_region))
    if missing_city:
        print("WARNING 城市BM Sheet 缺少行: " + ", ".join(missing_city))
    print(f"飞书写入完成。Klook新增周块: {klook_inserted}，城市BM新增周块: {city_inserted}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\n已停止。")
        raise SystemExit(130)
    except Exception as exc:
        print(f"ERROR {exc}", file=sys.stderr)
        raise SystemExit(1)
